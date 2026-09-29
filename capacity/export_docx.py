"""Export the capacity report to DOCX; install requirements-docx.txt first.

Supports the report's headings, paragraphs, bullets, inline emphasis/links,
tables and explicit page breaks. No network or service calls.
"""
from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/권장_서버_용량_추정.md"
OUTPUT = SOURCE.with_suffix(".docx")
INK = "4747B3"
FONT = "맑은 고딕"


def element(tag, **attrs):
    node = OxmlElement(f"w:{tag}")
    for key, value in attrs.items():
        node.set(qn(f"w:{key}"), str(value))
    return node


def inline(paragraph, text):
    pattern = r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\))"
    for part in re.split(pattern, text):
        if not part:
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", part)
        if link:
            label, url = link.groups()
            anchor = OxmlElement("w:hyperlink")
            anchor.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
            run = element("r")
            props = element("rPr")
            props.append(element("color", val=INK))
            props.append(element("u", val="single"))
            run.append(props)
            node = element("t")
            node.text = label
            run.append(node)
            anchor.append(run)
            paragraph._p.append(anchor)
        else:
            bold = part.startswith("**")
            code = part.startswith("`")
            run = paragraph.add_run(part[2:-2] if bold else part[1:-1] if code else part)
            run.bold = bold
            if code:
                run.font.size = Pt(9)


def add_table(doc, lines):
    rows = [[x.strip() for x in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
    if any(len(row) != len(rows[0]) for row in rows):
        raise ValueError("Inconsistent table columns")
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.autofit = False
    table.style = "Table Grid"
    # Wider narrative columns; scenario tables retain equal numeric columns.
    if rows[0][0] == "역할":
        widths = [3.0, 4.2, 4.2, 5.2]
    elif len(rows[0]) == 3:
        widths = [4.0, 5.0, 7.6]
    elif len(rows[0]) == 4:
        widths = [5.2, 3.8, 3.8, 3.8]
    else:
        widths = [16.6 / len(rows[0])] * len(rows[0])
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    for index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(element("cantSplit"))
        if index == 0:
            row._tr.get_or_add_trPr().append(element("tblHeader"))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.line_spacing = 1.1
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(9)
                if index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if index == 0 or index % 2 == 0:
                cell._tc.get_or_add_tcPr().append(element("shd", fill=INK if index == 0 else "F3F3FA"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    section.header_distance, section.footer_distance = Cm(.7), Cm(.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
        style = doc.styles[name]
        style.font.name = FONT
        props = style.element.get_or_add_rPr()
        props.rFonts.set(qn("w:eastAsia"), FONT)
        props.append(element("lang", val="ko-KR", eastAsia="ko-KR"))
    normal = doc.styles["Normal"]
    normal.font.size = Pt(9.5)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 25), ("Heading 1", 16), ("Heading 2", 12), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(7)
    header = section.header.paragraphs[0]
    header.add_run("G-HeSS AI Agent  |  서버 용량 계획").font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("2026-09-29  ·  계획 추정치  |  ").font.size = Pt(8)
    field = element("fldSimple", instr="PAGE")
    footer._p.append(field)
    doc.core_properties.title = "G-HeSS AI Agent — 트래픽 추정에 따른 권장 서버 용량"
    doc.core_properties.subject = "1차 오픈 트래픽, 서버 자원, 모델 쿼터 및 저장소 추정"
    doc.core_properties.author = "G-HeSS"
    doc.core_properties.language = "ko-KR"
    doc.core_properties.created = datetime(2026, 9, 29, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2026, 9, 29, tzinfo=timezone.utc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, table_lines)
            continue
        if line == "<!-- pagebreak -->":
            doc.add_page_break()
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], "Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2")
        elif line.startswith("- "):
            inline(doc.add_paragraph(style="List Bullet"), line[2:])
        elif line:
            inline(doc.add_paragraph(), line)
        index += 1
    doc.save(OUTPUT)
    print(f"Saved {OUTPUT.relative_to(ROOT)} ({OUTPUT.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
